import hashlib
import io
import os
import re
import shutil
import uuid
from pathlib import Path

from fastapi import UploadFile
from openpyxl import load_workbook
from PIL import Image
from docx import Document
import pikepdf

ALLOWED_EXTENSIONS = {"pdf", "jpg", "jpeg", "png", "docx", "xlsx"}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
UPLOAD_ROOT = Path("/app/uploads")


class UploadValidationError(ValueError):
    pass


def validate_and_store_upload(upload: UploadFile, tenant_id, report_id) -> dict:
    original_filename = upload.filename or "arquivo"
    extension = _extension(original_filename)
    if extension not in ALLOWED_EXTENSIONS:
        raise UploadValidationError("Tipo de arquivo nao permitido.")

    data = upload.file.read()
    if len(data) > MAX_UPLOAD_BYTES:
        raise UploadValidationError("Arquivo excede o limite de 10MB.")
    if not data:
        raise UploadValidationError("Arquivo vazio.")

    try:
        sanitized = sanitize_file(data, extension)
    except Exception as exc:
        raise UploadValidationError("Arquivo invalido ou corrompido.") from exc
    stored_filename = f"{uuid.uuid4()}.{extension}"
    directory = UPLOAD_ROOT / str(tenant_id) / str(report_id)
    directory.mkdir(parents=True, exist_ok=True)
    storage_path = directory / stored_filename
    storage_path.write_bytes(sanitized)

    return {
        "original_filename": _safe_name(original_filename),
        "stored_filename": stored_filename,
        "storage_path": str(storage_path),
        "mime_type": upload.content_type or "application/octet-stream",
        "size_bytes": len(sanitized),
        "sha256_hash": hashlib.sha256(sanitized).hexdigest(),
    }


def sanitize_file(data: bytes, extension: str) -> bytes:
    if extension in {"jpg", "jpeg", "png"}:
        return _sanitize_image(data, extension)
    if extension == "pdf":
        return _sanitize_pdf(data)
    if extension == "docx":
        return _sanitize_docx(data)
    if extension == "xlsx":
        return _sanitize_xlsx(data)
    raise UploadValidationError("Tipo de arquivo nao permitido.")


def _sanitize_image(data: bytes, extension: str) -> bytes:
    with Image.open(io.BytesIO(data)) as image:
        output = io.BytesIO()
        image = image.convert("RGB") if extension in {"jpg", "jpeg"} else image.copy()
        fmt = "JPEG" if extension in {"jpg", "jpeg"} else "PNG"
        image.save(output, format=fmt)
        return output.getvalue()


def _sanitize_pdf(data: bytes) -> bytes:
    output = io.BytesIO()
    with pikepdf.open(io.BytesIO(data)) as pdf:
        with pdf.open_metadata(set_pikepdf_as_editor=False) as meta:
            meta.clear()
        pdf.save(output)
    return output.getvalue()


def _sanitize_docx(data: bytes) -> bytes:
    src = io.BytesIO(data)
    document = Document(src)
    props = document.core_properties
    props.author = ""
    props.comments = ""
    props.keywords = ""
    props.last_modified_by = ""
    props.subject = ""
    props.title = ""
    output = io.BytesIO()
    document.save(output)
    # TODO: Sprint futura pode remover propriedades customizadas OOXML mais profundas.
    return output.getvalue()


def _sanitize_xlsx(data: bytes) -> bytes:
    workbook = load_workbook(io.BytesIO(data))
    props = workbook.properties
    props.creator = ""
    props.lastModifiedBy = ""
    props.title = ""
    props.subject = ""
    props.keywords = ""
    props.description = ""
    output = io.BytesIO()
    workbook.save(output)
    # TODO: Sprint futura pode remover propriedades customizadas OOXML mais profundas.
    return output.getvalue()


def _extension(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _safe_name(filename: str) -> str:
    name = os.path.basename(filename)
    return re.sub(r"[^A-Za-z0-9._-]", "_", name)[:255]
